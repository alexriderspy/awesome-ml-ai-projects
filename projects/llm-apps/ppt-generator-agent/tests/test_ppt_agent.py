import json
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "eval"))

from ppt_agent.build import build_pptx
from ppt_agent.plan import Deck, generate_deck, parse_json
from ppt_agent.render import render_with_pillow
from run_eval import bullet_compliance, coverage, pptx_content_slides, score

OUTLINE = {"title": "Docker 101", "subtitle": "Containers for developers",
           "sections": [{"title": "Basics", "slides": ["Images", "Containers"]},
                        {"title": "Shipping", "slides": ["Registries"]}]}


def slides_for(titles):
    return {"slides": [{"title": t, "bullets": [f"{t} fact one", "fact two", "fact three"],
                        "notes": "Say this.", "image_query": None} for t in titles]}


class FakeLLM:

    def __init__(self, replies):
        self.replies = list(replies)
        self.seen = []

    def chat(self, messages):
        self.seen.append(messages)
        return self.replies.pop(0)


def test_parse_json_strips_fences():
    assert parse_json('```json\n{"a": 1}\n```') == {"a": 1}
    assert parse_json('Sure! {"a": [1, 2]} hope that helps') == {"a": [1, 2]}
    with pytest.raises(ValueError):
        parse_json("no json here")


def test_generate_deck_happy_path():
    llm = FakeLLM([json.dumps(OUTLINE), json.dumps(slides_for(["Images", "Containers"])), json.dumps(slides_for(["Registries"]))])
    deck, stats = generate_deck("Docker", llm, n_slides=3)
    assert [s.title for s in deck.content_slides()] == ["Images", "Containers", "Registries"]
    assert stats.calls == 3 and stats.first_try_valid == 3 and stats.failed_steps == 0


def test_wrong_slide_count_gets_feedback_and_retry():
    bad = dict(OUTLINE, sections=[{"title": "Basics", "slides": ["Images"]}])
    llm = FakeLLM(["not json", json.dumps(bad), json.dumps(OUTLINE),
                   json.dumps(slides_for(["Images", "Containers"])), json.dumps(slides_for(["Registries"]))])
    deck, stats = generate_deck("Docker", llm, n_slides=3)
    assert len(deck.content_slides()) == 3
    assert stats.first_try_valid == 2
    feedback = llm.seen[2][-1]["content"]
    assert "exactly 3 are required" in feedback


def test_long_bullets_are_sent_back():
    long = slides_for(["Registries"])
    long["slides"][0]["bullets"][0] = " ".join(["word"] * 20)
    llm = FakeLLM([json.dumps(OUTLINE), json.dumps(slides_for(["Images", "Containers"])), json.dumps(long),
                   json.dumps(slides_for(["Registries"]))])
    deck, stats = generate_deck("Docker", llm, n_slides=3)
    assert "max is 12" in llm.seen[-1][-1]["content"]
    assert bullet_compliance(deck) == 1.0


def make_deck():
    llm = FakeLLM([json.dumps(OUTLINE), json.dumps(slides_for(["Images", "Containers"])), json.dumps(slides_for(["Registries"]))])
    return generate_deck("Docker", llm, n_slides=3)


def test_build_pptx_has_layouts_and_notes(tmp_path):
    from pptx import Presentation

    deck, _ = make_deck()
    out = build_pptx(deck, tmp_path / "d.pptx")
    prs = Presentation(str(out))
    assert len(prs.slides) == 6
    assert pptx_content_slides(out) == 3
    content = prs.slides[2]
    assert content.shapes.title.text == "Images"
    assert content.notes_slide.notes_text_frame.text == "Say this."


def test_build_with_image(tmp_path):
    from PIL import Image
    from pptx import Presentation

    img = tmp_path / "x.png"
    Image.new("RGB", (400, 300), "gray").save(img)
    deck, _ = make_deck()
    out = build_pptx(deck, tmp_path / "d.pptx", {0: {"path": img, "credit": "Someone / CC BY", "page": "http://x"}})
    slide = Presentation(str(out)).slides[2]
    assert any(s.shape_type == 13 for s in slide.shapes)
    assert "CC BY" in slide.notes_slide.notes_text_frame.text


def test_pillow_previews(tmp_path):
    deck, _ = make_deck()
    paths = render_with_pillow(deck, tmp_path, [0, 2])
    assert [p.name for p in paths] == ["slide-01.png", "slide-03.png"]


def test_scorer():
    deck, stats = make_deck()
    case = {"n_slides": 3, "keywords": [["image"], ["registr"], ["kubernetes"]]}
    assert coverage(deck, case["keywords"]) == pytest.approx(2 / 3)
    r = score(deck, stats, case, n_in_file=3)
    assert r["count"] == 1.0 and r["json"] == 1.0 and r["bullets"] == 1.0
    assert score(deck, stats, case, n_in_file=1)["score"] < r["score"]
    empty = Deck(title="x", sections=[])
    assert bullet_compliance(empty) == 0.0
