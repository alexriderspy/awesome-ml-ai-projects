from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

from .plan import Deck, Slide

INK = RGBColor(0x11, 0x11, 0x11)
MUTED = RGBColor(0x66, 0x66, 0x66)
ACCENT = RGBColor(0x2F, 0x5D, 0xA8)
FONT = "Arial"

W, H = Inches(13.333), Inches(7.5)
TITLE_LAYOUT, CONTENT_LAYOUT, SECTION_LAYOUT = 0, 1, 2


def _place(shape, left, top, width, height):
    shape.left, shape.top, shape.width, shape.height = Emu(left), Emu(top), Emu(width), Emu(height)


def _style(text_frame, size: int, color=INK, bold=False):
    text_frame.word_wrap = True
    for p in text_frame.paragraphs:
        for r in p.runs:
            r.font.size, r.font.name, r.font.bold = Pt(size), FONT, bold
            r.font.color.rgb = color


def _bar(slide, left, top, width, height=Inches(0.08), color=ACCENT):
    bar = slide.shapes.add_shape(1, left, top, width, height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()
    return bar


def _footer(slide, text: str):
    box = slide.shapes.add_textbox(Inches(0.6), H - Inches(0.55), W - Inches(1.2), Inches(0.35))
    box.text_frame.text = text
    _style(box.text_frame, 11, MUTED)


def _title_slide(prs, deck: Deck):
    s = prs.slides.add_slide(prs.slide_layouts[TITLE_LAYOUT])
    title, sub = s.placeholders[0], s.placeholders[1]
    _place(title, Inches(0.9), Inches(2.3), Inches(11.5), Inches(1.6))
    _place(sub, Inches(0.9), Inches(4.1), Inches(11.5), Inches(1.0))
    title.text, sub.text = deck.title, deck.subtitle
    for tf in (title.text_frame, sub.text_frame):
        for p in tf.paragraphs:
            p.alignment = 1
    _style(title.text_frame, 44, INK, bold=True)
    _style(sub.text_frame, 22, MUTED)
    _bar(s, Inches(0.9), Inches(2.1), Inches(1.4))


def _section_slide(prs, idx: int, title: str):
    s = prs.slides.add_slide(prs.slide_layouts[SECTION_LAYOUT])
    head, body = s.placeholders[0], s.placeholders[1]
    _place(body, Inches(0.9), Inches(2.2), Inches(4), Inches(0.8))
    _place(head, Inches(0.9), Inches(3.0), Inches(11.5), Inches(1.6))
    body.text, head.text = f"{idx:02d}", title
    _style(body.text_frame, 24, ACCENT, bold=True)
    _style(head.text_frame, 40, INK, bold=True)
    _bar(s, Inches(0.9), Inches(4.8), Inches(11.5), Inches(0.03), MUTED)


def _content_slide(prs, slide: Slide, footer: str, image: dict | None):
    s = prs.slides.add_slide(prs.slide_layouts[CONTENT_LAYOUT])
    head, body = s.placeholders[0], s.placeholders[1]
    _place(head, Inches(0.6), Inches(0.45), Inches(12.1), Inches(1.1))
    text_w = Inches(7.2) if image else Inches(12.1)
    _place(body, Inches(0.6), Inches(1.8), text_w, Inches(4.9))
    head.text = slide.title
    _style(head.text_frame, 32, INK, bold=True)
    head.text_frame.paragraphs[0].alignment = 1
    _bar(s, Inches(0.6), Inches(1.5), Inches(0.9))

    tf = body.text_frame
    for i, b in enumerate(slide.bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = b
        p.space_after = Pt(10)
    _style(tf, 22)

    if image:
        pic = s.shapes.add_picture(str(image["path"]), Inches(8.2), Inches(1.8), width=Inches(4.5))
        if pic.height > Inches(4.3):
            ratio = Inches(4.3) / pic.height
            pic.height, pic.width = Inches(4.3), int(pic.width * ratio)
        cap = s.shapes.add_textbox(Inches(8.2), pic.top + pic.height + Inches(0.05), Inches(4.5), Inches(0.4))
        cap.text_frame.text = image["credit"]
        _style(cap.text_frame, 9, MUTED)

    notes = slide.notes + (f"\n\nImage: {image['credit']} ({image['page']})" if image else "")
    s.notes_slide.notes_text_frame.text = notes.strip()
    _footer(s, footer)


def build_pptx(deck: Deck, out: str | Path, images: dict[int, dict] | None = None) -> Path:
    images = images or {}
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    _title_slide(prs, deck)
    n = 0
    for si, sec in enumerate(deck.sections, 1):
        _section_slide(prs, si, sec.title)
        for slide in sec.slides:
            _content_slide(prs, slide, f"{deck.title}  ·  {sec.title}", images.get(n))
            n += 1
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out
