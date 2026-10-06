import argparse
import json
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from ppt_agent.build import CONTENT_LAYOUT, build_pptx
from ppt_agent.plan import MAX_BULLET_WORDS, MAX_BULLETS, MIN_BULLETS, Deck, Stats, generate_deck

PRACTICE = Path(__file__).parent / "practice"


def pptx_content_slides(path: Path) -> int:
    from pptx import Presentation

    prs = Presentation(str(path))
    name = prs.slide_layouts[CONTENT_LAYOUT].name
    return sum(1 for s in prs.slides if s.slide_layout.name == name)


def bullet_compliance(deck: Deck) -> float:
    slides = deck.content_slides()
    if not slides:
        return 0.0
    ok = sum(
        1 for s in slides
        if MIN_BULLETS <= len(s.bullets) <= MAX_BULLETS and all(len(b.split()) <= MAX_BULLET_WORDS for b in s.bullets)
    )
    return ok / len(slides)


def deck_text(deck: Deck) -> str:
    parts = [deck.title, deck.subtitle]
    for sec in deck.sections:
        parts.append(sec.title)
        for s in sec.slides:
            parts += [s.title, *s.bullets, s.notes]
    return " " + " ".join(parts).lower() + " "


def coverage(deck: Deck, keywords: list[list[str]]) -> float:
    text = deck_text(deck)
    return sum(any(k.lower() in text for k in group) for group in keywords) / len(keywords)


def score(deck: Deck, stats: Stats, case: dict, n_in_file: int) -> dict:
    n = case["n_slides"]
    parts = {
        "json": stats.first_try_valid / stats.steps if stats.steps else 0.0,
        "count": 1.0 if n_in_file == n else max(0.0, 1 - abs(n_in_file - n) / n),
        "bullets": bullet_compliance(deck),
        "coverage": coverage(deck, case["keywords"]),
    }
    parts["score"] = round(100 * (0.2 * parts["json"] + 0.25 * parts["count"] + 0.25 * parts["bullets"] + 0.3 * parts["coverage"]), 1)
    return parts


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--json", action="store_true")
    args = ap.parse_args()

    from ppt_agent.llm import OpenAICompatible

    llm = OpenAICompatible()
    results = {}
    with tempfile.TemporaryDirectory() as tmp:
        for path in sorted(PRACTICE.glob("*.json")):
            case = json.loads(path.read_text())
            deck, stats = generate_deck(case["topic"], llm, n_slides=case["n_slides"])
            out = build_pptx(deck, Path(tmp) / f"{case['id']}.pptx")
            r = score(deck, stats, case, pptx_content_slides(out))
            r["llm_calls"] = stats.calls
            results[case["id"]] = r

    overall = round(sum(r["score"] for r in results.values()) / len(results), 1)
    if args.json:
        print(json.dumps({"overall": overall, "model": llm.model, "cases": results}, indent=2))
        return
    for cid, r in results.items():
        print(f"{cid:17s} score={r['score']:5.1f}  json={r['json']:.2f}  count={r['count']:.2f}  "
              f"bullets={r['bullets']:.2f}  coverage={r['coverage']:.2f}  calls={r['llm_calls']}")
    print(f"\nmodel: {llm.model}")
    print(f"OVERALL {overall}/100")


if __name__ == "__main__":
    main()
